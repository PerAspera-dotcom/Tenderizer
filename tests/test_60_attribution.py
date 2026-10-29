"""BOSA e-Procurement go-live conditions (2026-10-29): Belgian notices
carry attribution + a disclaimer wherever they leave the app by email."""
import attribution
import digest
import store
from conftest import TEST_TENANT_ID
from test_35_digest import TODAY, _tender


def test_email_footer_only_for_sources_with_a_notice():
    assert attribution.email_footer(["TED", "BOAMP"]) == ""
    footer = attribution.email_footer(["TED", "e-Procurement", "e-Procurement"])
    assert footer.count("publicprocurement.be") == 1
    assert "not affiliated with or endorsed by FPS BOSA" in footer


def test_digest_appends_bosa_notice_for_belgian_tender(tmp_path):
    conn = store.init_db(str(tmp_path / "t.db"))
    store.upsert(conn, TEST_TENANT_ID, {**_tender("BE-1"), "source": "e-Procurement"})
    body = digest.build_daily_digest(conn, TEST_TENANT_ID, today=TODAY)
    assert "Source: Belgian e-Procurement platform" in body


def test_digest_has_no_notice_without_belgian_tender(tmp_path):
    conn = store.init_db(str(tmp_path / "t.db"))
    store.upsert(conn, TEST_TENANT_ID, _tender("TED-1"))
    body = digest.build_daily_digest(conn, TEST_TENANT_ID, today=TODAY)
    assert "FPS BOSA" not in body


def test_report_adds_sources_sheet_only_with_belgian_tender(tmp_path):
    from openpyxl import load_workbook
    import report
    rec = {**_tender("BE-1"), "category": "Supply"}
    path = report.build_report([{**rec, "source": "e-Procurement"}], {}, str(tmp_path / "be.xlsx"))
    ws = load_workbook(path)["Sources"]
    assert "publicprocurement.be" in ws["A1"].value
    path = report.build_report([rec], {}, str(tmp_path / "ted.xlsx"))
    assert "Sources" not in load_workbook(path).sheetnames
